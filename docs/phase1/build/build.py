import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x00, 0x20, 0x60)
RED = RGBColor(0xC0, 0x00, 0x00)
BLANK = RGBColor(0xFF, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xED, 0xF6)
GREEN = RGBColor(0x1E, 0x7B, 0x34)
FONT = "Bookman Old Style"

prs = Presentation("stage.pptx")
S = list(prs.slides)


def body_ph(slide):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 1:
            return sh


def title_box(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip():
            r = sh.text_frame.paragraphs[0].runs
            if r and r[0].font.size and r[0].font.size.pt >= 30:
                return sh


def set_title(slide, text):
    tb = title_box(slide)
    tf = tb.text_frame
    paras = tf.paragraphs
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return tb


def remove(shape):
    shape._element.getparent().remove(shape._element)


def fill_body(slide, items, size=18, space_before=8, ph=None):
    """items: list of (runs, level) where runs = [(text, bold, color)] or a plain string."""
    ph = ph or body_ph(slide)
    txBody = ph.text_frame._txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    paras = txBody.findall(qn("a:p"))
    tmpl = copy.deepcopy(paras[0])
    for p in paras:
        txBody.remove(p)
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = ([(item, False, None)], 0)
        runs, lvl = item if isinstance(item, tuple) else (item, 0)
        p = copy.deepcopy(tmpl)
        for el in p.findall(qn("a:r")) + p.findall(qn("a:endParaRPr")) + p.findall(qn("a:br")):
            p.remove(el)
        pPr = p.find(qn("a:pPr"))
        if pPr is None:
            pPr = etree.SubElement(p, qn("a:pPr"))
            p.insert(0, pPr)
        if lvl == 1:
            pPr.set("lvl", "1"); pPr.set("marL", "571500"); pPr.set("indent", "-228600")
            bu = pPr.find(qn("a:buChar"))
            if bu is not None: bu.set("char", "–")
        if lvl == -1:  # no bullet
            pPr.set("marL", "0"); pPr.set("indent", "0")
            for tag in ("a:buChar", "a:buAutoNum", "a:buClr", "a:buSzPts", "a:buFont"):
                for el in pPr.findall(qn(tag)): pPr.remove(el)
            etree.SubElement(pPr, qn("a:buNone"))
        sb = pPr.find(qn("a:spcBef"))
        if sb is not None:
            sb.find(qn("a:spcPts")).set("val", str((space_before if i else 0) * 100))
        for text, bold, color in runs:
            r = etree.SubElement(p, qn("a:r"))
            rPr = etree.SubElement(r, qn("a:rPr"), lang="en-IN", sz=str(int((size - (2 if lvl == 1 else 0)) * 100)), b="1" if bold else "0")
            if color is not None:
                sf = etree.SubElement(rPr, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr"), val=str(color))
            t = etree.SubElement(r, qn("a:t")); t.text = text
            if text != text.strip():
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        txBody.append(p)
    return ph


def B(label, rest="", lvl=0, rest_color=None):
    runs = [(label, True, None)]
    if rest:
        runs.append((rest, False, rest_color))
    return (runs, lvl)


def P(text, lvl=0):
    return ([(text, False, None)], lvl)


def txt(slide, x, y, w, h, text, size=14, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        segs = line if isinstance(line, list) else [(line, bold, color)]
        for s, b, c in segs:
            r = p.add_run(); r.text = s
            f = r.font; f.size = Pt(size); f.bold = b; f.italic = italic; f.name = FONT; f.color.rgb = c or color
    return tb


def box(slide, x, y, w, h, text, fill=NAVY, fg=WHITE, size=13, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        f = r.font; f.size = Pt(size if i == 0 else size - 2); f.bold = bold if i == 0 else False; f.name = FONT; f.color.rgb = fg
    return s


def arrow(slide, x1, y1, x2, y2, color=NAVY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def table(slide, x, y, w, col_w, rows, size=12, header_size=13, row_h=0.36):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(0.45 if i == 0 else row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            tf = cell.text_frame; tf.word_wrap = True
            segs = val if isinstance(val, list) else [(val, i == 0 or j == 0, None)]
            p = tf.paragraphs[0]
            for s, b, c in segs:
                r = p.add_run(); r.text = s
                f = r.font; f.name = FONT; f.size = Pt(header_size if i == 0 else size); f.bold = b
                f.color.rgb = WHITE if i == 0 else (c or NAVY)
    return t


# ---------- 1. Title ----------
s = S[0]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "THE MEDMINDS":
        sh.text_frame.paragraphs[0].runs[0].text = "OPTIFINERS"
    elif sh.has_text_frame and sh.text_frame.text.strip() == "BIOMEDICAL AND HEALTHCARE":
        sh.text_frame.paragraphs[0].runs[0].text = "INDUSTRIAL ENGINEERING & MANAGEMENT"; sh.text_frame.paragraphs[0].runs[0].font.size = Pt(32)

# ---------- 2. Tentative steps ----------
s = S[1]
remove(body_ph(s))
txt(s, 2.3, 0.2, 10.6, 0.8, "Outline of Tentative Steps", size=36, bold=True, color=RED, align=PP_ALIGN.CENTER)
txt(s, 2.3, 1.0, 10.6, 0.45, "OptiProductio: Smart Production Workflow and Resource Optimization System", size=16, bold=True, align=PP_ALIGN.CENTER)
phases = [
    ("Phase I", "29 Oct 2026", ["Pick and define the problem", "Literature review", "SMART objectives (PO1 to PO4)", "Design the 3-station line and parts list", "Plan the method and timeline"]),
    ("Phase II", "__ / __ / 2026", ["Wire sensors to Arduino UNO", "Python logger and dashboard", "Check counts against a hand tally", "First runs and partial results", "Dashboard on Bosch public data"]),
    ("Phase III", "__ / __ / 2027", ["Rebalance and WIP-limit experiments", "Results and discussion", "Live prototype demo", "Report, poster and PPT", "IEEE-format paper draft"]),
]
for k, (name, date, steps) in enumerate(phases):
    x = 0.55 + k * 4.15
    b = box(s, x, 1.7, 3.85, 0.85, [name, date], size=20)
    if "__" in date:
        b.text_frame.paragraphs[1].runs[0].font.color.rgb = RGBColor(0xFF, 0x9B, 0x9B)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.7), Inches(3.85), Inches(3.55))
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.fill.background(); card.shadow.inherit = False
    card.adjustments[0] = 0.05
    txt(s, x + 0.15, 2.8, 3.6, 3.4, [[("• " + st, False, NAVY)] for st in steps], size=17)
    if k < 2:
        arrow(s, x + 3.88, 2.12, x + 4.12, 2.12)
txt(s, 0.55, 6.35, 12.3, 0.4, [[("Status today: Phase I planning done. Hardware build in progress. Red blanks = dates not announced yet.", False, NAVY)]], size=12, italic=True)

# ---------- 3. Outline ----------
s = S[2]
tb = title_box(s)
fill_body(s, [P(x) for x in ["Introduction", "Literature Review", "Problem Definition", "Objectives", "Methodology and System Pipeline",
                             "Measurement Parameters", "Research Process", "Proposed Outcomes", "Timeline of the Project", "Conclusion", "References",
                             "Expected Outcome by the End of First Year"]], size=20, space_before=3)

# ---------- 4. Introduction ----------
s = S[3]
set_title(s, "INTRODUCTION")
ph = body_ph(s)
ph.top = Inches(1.15); ph.height = Inches(3.2)
fill_body(s, [
    ([("OVERVIEW OF THE TOPIC:", True, None)], -1),
    B("OptiProductio: ", "a low-cost system that watches a small production line in real time and shows how well it is working."),
    B("How: ", "sensors on a 3-station model line send every event (part done, part rejected, line stopped) to an Arduino UNO, which passes it to a Python dashboard on a laptop."),
    B("Why: ", "small factories often record output by hand, so delays, pile-ups and the slowest station stay hidden until the end of the shift [2], [4]."),
    B("What it adds: ", "an alert when too many parts pile up inside the line (a WIP limit), based on the CONWIP idea [5]."),
], size=17, space_before=8)
labels = [("Production count", "IR sensor at each station exit"), ("Cycle time", "time each station takes per part"),
          ("Inventory", "raw material bin level and WIP"), ("Workflow", "bottleneck, stops and pile-ups")]
for k, (a, b2) in enumerate(labels):
    box(s, 0.45 + k * 3.15, 4.6, 2.95, 1.25, [a, b2], size=16)
txt(s, 0.45, 6.0, 12.4, 0.4, "EL topic: Smart Production Workflow and Resource Optimization System. Deliverable: production efficiency dashboard.", size=12, italic=True)

# ---------- 5. Literature review (key studies) ----------
s = S[4]
set_title(s, "Literature Review")
remove(body_ph(s))
txt(s, 0.4, 1.12, 12.5, 0.4, "Key studies and findings", size=18, bold=True)
rows = [
    ["Study", "What they did", "Key finding"],
    ["Hwang et al., 2017 [2]", "IoT performance-measurement model built on the ISA-95 and ISO 22400 standards", "IoT timestamp data can give OEE numbers in real time"],
    ["Ng Corrales et al., 2020 [3]", "Systematic review of OEE: 862 articles screened down to 186", "Many industries adapt OEE, because the basic formula alone is not enough"],
    ["Wicaksono et al., 2024 [4]", "Replaced manual OEE logging with an ESP32 board over MQTT", "Works in real time, but reports up to 16% OEE error"],
    ["Spearman et al., 1990 [5]", "Proposed CONWIP: a fixed cap on the total parts inside a line", "A simple pull alternative to kanban that limits WIP"],
    ["Thürer et al., 2023 [6]", "Tested kanban, CONWIP and a hybrid on a physical lab-scale line", "WIP limits cut the time a job spends in the line but can lower throughput"],
]
table(s, 0.4, 1.6, 12.5, [2.7, 5.1, 4.7], rows, size=14, header_size=15, row_h=0.8)

# ---------- 6. Literature review (contd.) ----------
s = S[5]
set_title(s, "Literature Review (contd.)")
fill_body(s, [
    B("Trends in literature: ", "low-cost IoT boards now give real-time OEE [2], [4]; WIP limits like CONWIP are mostly tested in simulation or lab models [6], [9]; public line data such as the Bosch dataset is used mainly to predict failed parts [8]."),
    B("Theoretical framework:", ""),
    B("Little's law: ", "average WIP = throughput × average flow time [7]. Cut WIP at the same throughput and parts spend less time in the line.", 1),
    B("OEE = Availability × Performance × Quality ", "[3], [10].", 1),
    B("CONWIP: ", "release a new part only when a finished part leaves [5].", 1),
    B("Relevance to our project: ", "monitoring papers report numbers but do not control the flow, and WIP-limit papers control the flow but are rarely tied to a live, measured line [6], [9]. Low-cost builds rarely check their accuracy against a manual count [4]. OptiProductio does all three on one small line."),
], size=17, space_before=10)

# ---------- 7. Problem definition ----------
s = S[6]
set_title(s, "Problem Definition")
fill_body(s, [
    B("Clear statement of the problem: ", "small production units cannot see count, cycle time, inventory and WIP while the line is running, so they find bottlenecks and pile-ups too late to act."),
    B("Context of the problem: ", "MSMEs gave 30.1% of India's GDP (2022-23) and 45.73% of its exports (2023-24) [1]. Manual OEE recording is slow because operators write data at the machine and supervisors process it later [4]. Ordinary performance systems struggle to capture what is actually happening on the floor in real time [2]."),
    B("Significance of addressing it: ", "if work is released without a limit, WIP between stations can keep growing [11], parts wait longer, and the slowest station stays hidden."),
    B("Gap and novelty: ", "low-cost monitoring and WIP control are studied separately [4], [9]. OptiProductio puts both on one prototype costing under ₹2,000 (our estimate), and checks its own counts against a hand tally and Little's law [7]."),
], size=18, space_before=12)

# ---------- 8. Objectives ----------
s = S[7]
set_title(s, "Objectives")
ph = body_ph(s); ph.top = Inches(1.1); ph.height = Inches(1.2)
fill_body(s, [
    B("Main objective: ", "build a low-cost hardware and software prototype that measures production count, cycle time, inventory and WIP on a 3-station model line in real time, and use it to find the bottleneck and test a WIP limit."),
], size=17)
rows = [
    ["PO", "Specific objective", "Measurable target", "By"],
    ["PO1", "Count parts at all 3 stations with IR sensors", "Count error ≤ 2% against a hand tally, over three 10-minute runs", "Phase II"],
    ["PO2", "Live dashboard: count, cycle time per station, WIP, bin level, downtime, OEE proxy", "Refreshes every 2 s or faster; all 6 numbers on one screen", "Phase II"],
    ["PO3", "Find the bottleneck and rebalance the line", "Report throughput and flow time before and after, 3 runs each; Little's law gap ≤ 10%", "Phase III"],
    ["PO4", "WIP-limit alert (LED, buzzer, dashboard)", "Alert fires at the cap every time; trade-off reported for caps 2, 4, 6", "Phase III"],
]
t = table(s, 0.4, 2.3, 12.5, [0.8, 4.6, 5.4, 1.7], rows, size=14, header_size=15, row_h=0.62)
txt(s, 0.4, 5.6, 12.5, 0.9, [[("Why achievable: ", True, NAVY), ("parts are standard IDEA Lab kit, the code is already written and tested on sample data, and each PO is checked by a simple count or timer.  ", False, NAVY), ("Expected contribution: ", True, NAVY), ("a cheap, open template for monitoring and WIP control that small units and student labs can copy.", False, NAVY)]], size=15)

# ---------- 9. Methodology ----------
s = S[8]
set_title(s, "Methodology")
fill_body(s, [
    B("Design: ", "quantitative and experimental. Build the line, measure it, change one thing at a time (rebalance, then WIP limit) and compare the numbers."),
    B("Data collection: ", "IR sensors time every part leaving each station; buttons log breakdown, no-material and reject; an ultrasonic sensor reads the raw material bin. The Arduino stamps each event in milliseconds and sends it over USB to a Python logger (CSV). Secondary data: the public Bosch production line dataset [8], [12]."),
    B("Data analysis: ", "throughput, cycle time per station, WIP and flow time from the event log; Little's law as a cross-check [7]; OEE proxy = A × P × Q [3]; mean and spread across 3 repeat runs."),
    B("Tools: ", "Arduino IDE (C++), Python (pyserial, pandas, Streamlit), Excel for result tables."),
    B("Risks and fallbacks: ", "miscounts (cardboard gates, debounce, hand tally check); hardware failure on demo day (replay a saved run); Bosch time units are not published (use only for relative comparison)."),
], size=17, space_before=10)

# ---------- 10. System pipeline ----------
s = S[9]
set_title(s, "System Pipeline")
remove(body_ph(s))
# physical line row
y = 1.55
row = [("Raw material bin", "ultrasonic sensor"), ("Station 1: Cut", "IR sensor D2"), ("Station 2: Drill", "IR sensor D3 (slow on purpose)"), ("Station 3: Inspect", "IR D4 + Reject button")]
for k, (a, b2) in enumerate(row):
    x = 0.45 + k * 3.2
    box(s, x, y, 2.75, 1.0, [a, b2], fill=LIGHT, fg=NAVY, size=15, line=NAVY)
    if k < 3:
        arrow(s, x + 2.78, y + 0.5, x + 3.17, y + 0.5)
txt(s, 0.45, 1.12, 6, 0.35, "Physical line (parts moved by hand)", size=13, bold=True, color=RED)
box(s, 0.45, 3.00, 3.6, 1.0, ["Control panel", "Breakdown, No-material buttons"], fill=LIGHT, fg=NAVY, size=15, line=NAVY)
box(s, 4.85, 3.00, 3.6, 1.0, ["Arduino UNO", "timestamps every event"], size=16)
box(s, 9.25, 3.00, 3.6, 1.0, ["Andon on the line", "LCD, red/green LED, buzzer"], fill=LIGHT, fg=NAVY, size=15, line=NAVY)
arrow(s, 6.65, 2.55, 6.65, 2.98)
arrow(s, 4.07, 3.50, 4.83, 3.50)
arrow(s, 8.47, 3.50, 9.23, 3.50)
box(s, 4.85, 4.50, 3.6, 0.9, ["Python logger", "USB serial to CSV event log"], size=15)
arrow(s, 6.65, 4.02, 6.65, 4.48)
box(s, 0.45, 4.50, 3.6, 0.9, ["Bosch public dataset", "converted to the same events"], fill=LIGHT, fg=NAVY, size=15, line=NAVY)
arrow(s, 4.07, 4.95, 4.83, 4.95)
box(s, 4.85, 5.75, 3.6, 0.9, ["Metrics engine", "count, cycle time, WIP, OEE"], size=15)
arrow(s, 6.65, 5.42, 6.65, 5.73)
box(s, 9.25, 5.75, 3.6, 0.9, ["Streamlit dashboard", "live KPIs and WIP alert"], fill=RED, size=15)
arrow(s, 8.47, 6.20, 9.23, 6.20)
c = arrow(s, 11.05, 5.73, 11.05, 4.02, color=RED)
txt(s, 11.15, 4.60, 1.9, 0.6, "WIP cap sent back to Arduino", size=11, color=RED, italic=True)

# ---------- 11. Measurement parameters ----------
s = S[10]
set_title(s, "Measurement Parameters")
remove(body_ph(s))
txt(s, 0.4, 1.05, 12.5, 0.4, "No machine-learning model is trained in this project. These are the measurement settings we will report.", size=14, italic=True)
rows = [
    ["Item", "Value"],
    ["Data sources", "Live line event log (CSV) and a Bosch sample: first 50,000 rows, line L3, 3 busiest stations [12]"],
    ["Event format", "time (s), event (PASS / SCRAP / DOWN / UP), station, reason"],
    ["Sensor timing", "IR debounce 400 ms; button debounce 200 ms; status sent every 1 s; dashboard refresh 2 s"],
    ["Cycle time (per station)", "finish(k) − max(arrive(k), finish(k−1)), so stations after the bottleneck do not look slow"],
    ["WIP", "parts released at Station 1 minus parts that left Station 3 (good or scrap)"],
    ["Throughput / flow time", "good parts per minute / exit time − entry time per part; checked with Little's law [7]"],
    ["OEE proxy", [("A × P × Q [3]; ideal cycle time set by the team: ", False, None), ("__ s", True, BLANK)]],
    ["Experiments", "baseline, rebalanced, WIP cap 2 / 4 / 6; 10 minutes per run, 3 repeats each"],
    ["Targets", "count error ≤ 2%; Little's law gap ≤ 10%; alert fires every time WIP reaches the cap"],
]
table(s, 0.4, 1.5, 12.5, [3.2, 9.3], rows, size=14, header_size=15, row_h=0.5)

# ---------- 12. Research process ----------
s = S[11]
set_title(s, "Research Process")
remove(body_ph(s))
steps = ["Literature review", "Define problem and POs", "Design line and parts list", "Build and wire", "Calibrate: count vs hand tally"]
for k, st in enumerate(steps):
    x = 0.35 + k * 2.55
    box(s, x, 1.3, 2.2, 0.95, st, size=14)
    if k < 4:
        arrow(s, x + 2.22, 1.77, x + 2.53, 1.77)
d = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(10.15), Inches(2.6), Inches(2.6), Inches(1.4))
d.fill.solid(); d.fill.fore_color.rgb = WHITE; d.line.color.rgb = RED; d.line.width = Pt(2); d.shadow.inherit = False
tf = d.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run(); r.text = "Error ≤ 2%?"; r.font.size = Pt(14); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = RED
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
arrow(s, 11.45, 2.27, 11.45, 2.58)
txt(s, 8.6, 2.75, 1.5, 0.4, "No: adjust", size=12, color=RED, bold=True, align=PP_ALIGN.RIGHT)
arrow(s, 10.13, 3.3, 9.45, 3.3, color=RED)
box(s, 7.25, 2.85, 2.2, 0.9, "Fix gates and sensor range", fill=LIGHT, fg=NAVY, size=13, line=RED)
c = arrow(s, 8.35, 2.83, 8.35, 2.27, color=RED)
txt(s, 11.55, 4.0, 1.2, 0.4, "Yes", size=12, color=GREEN, bold=True)
steps2 = ["Run 1: baseline", "Find bottleneck", "Run 2: rebalanced", "Run 3: WIP caps 2/4/6", "Analyse: Little's law, OEE"]
for k, st in enumerate(steps2):
    x = 10.15 - k * 2.45
    box(s, x, 4.45, 2.2, 0.95, st, size=14, fill=NAVY if k != 3 else RED)
    if k < 4:
        arrow(s, x - 0.02, 4.92, x - 0.23, 4.92)
arrow(s, 11.45, 4.02, 11.45, 4.43)
box(s, 0.35, 5.85, 4.0, 0.8, "Repeat on Bosch public data", fill=LIGHT, fg=NAVY, size=14, line=NAVY)
box(s, 4.75, 5.85, 4.0, 0.8, "Report, poster, paper draft", fill=LIGHT, fg=NAVY, size=14, line=NAVY)
arrow(s, 1.45, 5.42, 1.45, 5.83)
arrow(s, 4.37, 6.25, 4.73, 6.25)

# ---------- 13. Proposed outcomes ----------
s = S[12]
set_title(s, "Proposed Outcomes")
fill_body(s, [
    B("Expected results: ", "a working prototype (3-station line and live dashboard) and a results table comparing baseline, rebalanced and WIP-capped runs: throughput, average WIP, flow time, cycle time per station, OEE proxy and scrap rate."),
    ([("Success criteria: ", True, None), ("PO1 to PO4 met. Improvement after rebalancing: ", False, None), ("__ %", True, BLANK), (" (to be measured, not assumed).", False, None)], 0),
    B("Potential impact: ", "shows that count, cycle time, inventory and workflow can be watched live with low-cost parts, as a template for small units and student labs."),
    B("Limits we will state: ", "OEE is a proxy (our own ideal cycle time); parts are moved by hand; Bosch time units are estimated by the Kaggle community, not published by Bosch [8]."),
    B("Future research directions: ", "wireless boards, more stations and product types, and testing the WIP-limit alert on more real factory data."),
], size=17, space_before=12)

# ---------- 14. Timeline ----------
s = S[13]
set_title(s, "Timeline of the Project")
remove(body_ph(s))
rows = [
    ["Sprint", "Dates", "Work", "Linked to", "Status"],
    ["1", "23 to 29 Oct 2026", "Problem, literature review, POs, line design, parts list, code skeleton", "Phase I review (29 Oct)", [("Done", True, GREEN)]],
    ["2", "30 Oct to 1 Nov 2026", "Wire all sensors, logger and dashboard running together", "PO1, PO2", [("In progress", True, RED)]],
    ["3", [("2 Nov to ", False, None), ("__", True, BLANK)], "Count accuracy test, Bosch data mode, first runs", "PO1, PO2, Phase II", [("Not started", False, None)]],
    ["4", [("__", True, BLANK), (" to ", False, None), ("__", True, BLANK)], "Baseline, rebalance and WIP-cap experiments (3 repeats each)", "PO3, PO4", [("Not started", False, None)]],
    ["5", [("__", True, BLANK), (" to ", False, None), ("__", True, BLANK)], "Results, report, poster, PPT, paper draft, demo rehearsal", "Phase III", [("Not started", False, None)]],
]
table(s, 0.4, 1.2, 12.5, [1.0, 2.6, 4.9, 2.3, 1.7], rows, size=15, header_size=16, row_h=0.75)
txt(s, 0.4, 5.75, 12.5, 0.7, [[("Phase II date: ", True, NAVY), ("__ / __ / 2026", True, BLANK), ("     Phase III date: ", True, NAVY), ("__ / __ / 2027", True, BLANK), ("     (dates to be announced; sprints will be fixed backwards from them)", False, NAVY)]], size=13)

# ---------- 15. Conclusion ----------
s = S[14]
set_title(s, "Conclusion")
fill_body(s, [
    B("Summary of key points: ", "OptiProductio measures production count, cycle time, inventory and workflow on a 3-station model line, live, using an Arduino UNO, low-cost sensors and a Python dashboard."),
    B("What makes it different: ", "it adds a WIP-limit alert based on CONWIP [5], and it checks its own numbers against a hand tally and Little's law [7]."),
    B("Final thought: ", "you cannot improve a line you cannot see. OptiProductio makes the line visible cheaply."),
    B("Next steps: ", "finish wiring by 1 Nov, run the count accuracy test (PO1), then bring partial results to Phase II."),
], size=20, space_before=16)

# ---------- 16. References ----------
s = S[15]
set_title(s, "References")
refs = [
    "[1] Press Information Bureau. (2024, December 23). The MSME revolution: Transforming India's economic landscape. Government of India. https://static.pib.gov.in/WriteReadData/specificdocs/documents/2024/dec/doc20241223476301.pdf",
    "[2] Hwang, G., Lee, J., Park, J., & Chang, T.-W. (2017). Developing performance measurement system for Internet of Things and smart factory environment. International Journal of Production Research, 55(9), 2590–2602. https://doi.org/10.1080/00207543.2016.1245883",
    "[3] Ng Corrales, L. del C., Lambán, M. P., Hernandez Korner, M. E., & Royo, J. (2020). Overall equipment effectiveness: Systematic literature review and overview of different approaches. Applied Sciences, 10(18), 6469. https://doi.org/10.3390/app10186469",
    "[4] Wicaksono, A. A., Ningsih, Y. K., & Surjati, I. (2024). Implementation of MQTT protocol on ESP32-based OEE analysis development board. Emitor: Jurnal Teknik Elektro, 24(2), 169–175. https://doi.org/10.23917/emitor.v24i2.3908",
    "[5] Spearman, M. L., Woodruff, D. L., & Hopp, W. J. (1990). CONWIP: A pull alternative to kanban. International Journal of Production Research, 28(5), 879–894. https://doi.org/10.1080/00207549008942761",
    "[6] Thürer, M., Li, S. S., Yang, C., Qu, T., & Huang, G. Q. (2023). Does regulating work-in-process increase throughput and reduce cycle times? An assessment by lab scale system models. In Advances in Production Management Systems (APMS 2023), IFIP AICT 691 (pp. 649–661). Springer. https://doi.org/10.1007/978-3-031-43670-3_45",
    "[7] Little, J. D. C. (2011). Little's law as viewed on its 50th anniversary. Operations Research, 59(3), 536–549. https://doi.org/10.1287/opre.1110.0940",
]
refs2 = [
    "[8] Mangal, A., & Kumar, N. (2017). Using big data to enhance the Bosch production line performance: A Kaggle challenge. arXiv. https://arxiv.org/abs/1701.00705",
    "[9] Pettersen, J.-A., & Segerstedt, A. (2009). Restricted work-in-process: A study of differences between Kanban and CONWIP. International Journal of Production Economics, 118(1), 199–207.",
    "[10] International Organization for Standardization. (2014). Automation systems and integration: Key performance indicators (KPIs) for manufacturing operations management. Part 2 (ISO 22400-2:2014). https://www.iso.org/standard/54497.html",
    "[11] Cordon, C. (1995). Quality defaults and work-in-process inventory. European Journal of Operational Research, 80(2), 240–251.",
    "[12] Bosch. (2016). Bosch production line performance [Data set]. Kaggle. https://www.kaggle.com/c/bosch-production-line-performance",
    "[13] The Patents Act, 1970, § 3(k) (India). https://www.indiacode.nic.in/handle/123456789/1392",
]
ph = body_ph(s); ph.top = Inches(1.15); ph.height = Inches(5.7)
fill_body(s, [([(r, False, None)], -1) for r in refs], size=15, space_before=8)

# ---------- 16b. References contd ----------
s = S[16]
set_title(s, "References (contd.)")
ph = body_ph(s); ph.top = Inches(1.15); ph.height = Inches(5.7)
fill_body(s, [([(r, False, None)], -1) for r in refs2], size=15, space_before=8)

# ---------- 17. Expected outcome ----------
s = S[17]
set_title(s, "Expected Outcome by the End of First Year")
fill_body(s, [
    B("Product prototype: ", "a working OptiProductio line and dashboard, released as open source (code, wiring table and build guide) so other labs can rebuild it."),
    B("Conference publication: ", "a short paper in IEEE format on the prototype, the accuracy check and the WIP-limit experiments. Target venue to be chosen once results are in."),
    B("Research proposal: ", "extend the work to more stations, wireless boards and more real factory data."),
    B("Patent: ", "not planned. Software and algorithms on their own are generally not patentable in India (Patents Act, 1970, Section 3(k)) [13], so we will not claim one."),
], size=18, space_before=14)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 10:
        ph.text_frame.paragraphs[0].runs[0].text = "Thursday, 29 October 2026" if ph.text_frame.paragraphs[0].runs else None

prs.save("stage2.pptx")

# references continuation slide: add a second references slide via duplication later
import json
json.dump(refs2, open("refs2.json", "w"))
print("ok")
